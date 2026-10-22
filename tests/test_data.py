from __future__ import annotations

import csv
import json
import re
import shutil
from decimal import Decimal
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
import pyarrow.parquet as pq
import pytest

from traceatlas.importer import (
    CancelledError,
    DataError,
    WorkbookReader,
    inspect_workbook,
    parse_money,
    validate_import,
)
from traceatlas.jobs import execute_job
from traceatlas.query import QueryService


class LocalProject:
    """Exercise production file/query contracts without coupling to project transaction tests."""
    def __init__(self, path):
        self.path = path
        self.path.mkdir(parents=True, exist_ok=True)
        self._objects = {}

    def get_object(self, identifier):
        return self._objects[identifier]

    def objects(self, kind=None):
        return [item for item in self._objects.values() if kind is None or item['kind'] == kind]

    def apply(self, staging, manifest):
        for item in manifest['files']:
            target = self.path / item['destination']
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(staging / item['source'], target)
        for obj in manifest['objects']:
            self._objects[obj['id']] = {**obj, 'state': 'ready'}


def workbook(path, rows):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = '流水'
    sheet.append(['说明行'])
    sheet.append(['账号', '金额', '时间', '对方账户', '原文'])
    for row in rows:
        sheet.append(row)
    book.save(path)
    return path


def parameters(path, **extra):
    result = {'path': str(path), 'sheet': '流水', 'header_row': 2,
              'mapping': {'owner_account': 1, 'amount': 2, 'occurred_at': 3, 'counterparty_account': 4},
              'platform': 'bank', 'account_namespace': 'account'}
    result.update(extra)
    return result


def run_job(project, feature, params, input_objects=None):
    number = len(list(project.path.glob('staging*')))
    staging = project.path / f'staging{number}'
    spec = {'task_id': f'task{number}', 'attempt_id': f'attempt{number}', 'project_id': 'project',
            'project_path': str(project.path), 'feature_id': feature, 'module_version': '1.0.0',
            'parameters': params, 'inputs': input_objects or [],
            'input_dataset_versions': [obj['id'] for obj in (input_objects or [])]}
    result = execute_job(spec, staging, lambda *args: None, lambda: False)
    project.apply(staging, result)
    return result, staging


def test_money_never_rounds_or_uses_float():
    assert parse_money('99999999999999999999999999.123456789012') == Decimal('99999999999999999999999999.123456789012')
    assert parse_money('0.1') + parse_money('0.2') == Decimal('0.3')
    for value in ('oops', 'NaN', '1,234.56', '0.1234567890123', '100000000000000000000000000'):
        with pytest.raises(ValueError):
            parse_money(value)


def test_preview_and_partial_import_keep_source_rows_and_raw_fields(tmp_path):
    source = workbook(tmp_path / '输入.xlsx', [
        ['00123', '12345678901234567890.123456789012', '2026-10-22', '00045', '=1+1'],
        ['6222021234567890123', 'bad', '2026-10-22 12:13:14', None, 'NA'],
        ['00123', '-0.01', None, 'ABC', ''],
    ])
    preview = inspect_workbook(source, header_row=2)
    assert preview['row_numbers'] == [3, 4, 5]
    assert preview['rows'][0][0] == '00123'
    report = validate_import(parameters(source))
    assert report['input_count'] == 3
    assert report['rejected_count'] == 1
    assert report['errors'][0]['row'] == 4
    project = LocalProject(tmp_path / 'project')
    with pytest.raises(DataError) as error:
        run_job(project, 'xlsx.import', parameters(source))
    assert error.value.code == 'IMPORT_VALIDATION'
    assert project.objects() == []
    manifest, _ = run_job(project, 'xlsx.import', parameters(source, allow_partial=True))
    dataset = project.objects('dataset')[0]
    page = QueryService(project).page(dataset['id'], sort='source_row')
    assert page['total'] == 2
    assert [row['source_row'] for row in page['rows']] == [3, 5]
    first = page['rows'][0]
    assert first['owner_account'] == '00123'
    assert first['amount'] == Decimal('12345678901234567890.123456789012')
    assert first['time_precision'] == 'day'
    assert first['occurred_at'] == '2026-10-22'
    assert first['currency'] == 'unknown'
    assert first['timezone'] == 'unknown'
    raw = json.loads(first['raw_cells'])
    assert raw['5']['formula'] == '=1+1'
    assert raw['5']['value'] is None
    assert any(item['field'] == 'column_5' for item in json.loads(first['quality_flags']))
    source.unlink()
    preview = QueryService(project).source_preview(first['source_id'], '流水', 3)
    assert any(row['row_number'] == 3 and row['column_1'] == '00123' for row in preview['rows'])
    assert manifest['summary']['completion'] == 'partial'


def test_account_result_all_contributors_filters_exports_and_reopen(tmp_path):
    source = workbook(tmp_path / '输入.xlsx', [
        [f'{i:05d}', str(i) + '.123456789012', '2026-10-22 23:59:59', '00000', 'x']
        for i in range(100)
    ])
    project = LocalProject(tmp_path / '项目 中文')
    run_job(project, 'xlsx.import', parameters(source))
    dataset = project.objects('dataset')[0]
    result_manifest, _ = run_job(project, 'validation.account_count', {}, [dataset])
    result = project.objects('result')[0]
    assert result_manifest['summary']['contributing_count'] == 100
    query = QueryService(project)
    filters = [{'column': 'owner_account', 'op': 'lte', 'value': '00011'}]
    page = query.page(result['id'], filters, sort='owner_account')
    assert page['total'] == 12
    assert sum(query.sources(result['id'], row['result_row_id'])['total'] for row in page['rows']) == 12
    assert query.page(dataset['id'], [{'column': 'occurred_at', 'op': 'lte', 'value': '2026-10-22'}])['total'] == 100
    for scope, expected in [('filtered', 12), ('all', 100)]:
        destination = tmp_path / f'export-{scope}.csv'
        manifest, staging = run_job(project, 'result.export', {
            'object_id': result['id'], 'filters': filters, 'sort': 'owner_account', 'scope': scope,
            'format': 'csv', 'destination': str(destination),
        }, [result])
        with (staging / 'export.csv').open(encoding='utf-8-sig', newline='') as stream:
            rows = list(csv.DictReader(stream))
        assert len(rows) == expected
        assert rows[0]['owner_account'] == '00000'
        assert manifest['summary']['row_count'] == expected
        metadata = json.loads((staging / 'export.metadata.json').read_text(encoding='utf-8'))
        assert metadata['dataset_versions'] == [dataset['id']]
    _, staging = run_job(project, 'result.export', {
        'object_id': dataset['id'], 'format': 'xlsx', 'destination': str(tmp_path / 'exact.xlsx'),
        'scope': 'all', 'sort': 'source_row',
    }, [dataset])
    exported = openpyxl.load_workbook(staging / 'export.xlsx', read_only=True, data_only=False)
    headers, first = list(exported.active.values)[:2]
    assert first[headers.index('owner_account')] == '00000'
    assert first[headers.index('amount')] == '0.123456789012'
    exported.close()
    # Move and reopen the same registered relative references, rather than recomputing.
    moved = tmp_path / 'moved'
    shutil.move(project.path, moved)
    project.path = moved
    assert QueryService(project).page(result['id'])['total'] == 100
    row = QueryService(project).page(result['id'], sort='owner_account')['rows'][0]
    assert QueryService(project).sources(result['id'], row['result_row_id'])['total'] == 1


def test_count_keeps_duplicates_namespaces_and_multi_source_lineage(tmp_path):
    source = workbook(tmp_path / 'a.xlsx', [['00123', '1', None, None, None]] * 3)
    project = LocalProject(tmp_path / 'p')
    run_job(project, 'xlsx.import', parameters(source))
    run_job(project, 'xlsx.import', parameters(source))
    run_job(project, 'xlsx.import', parameters(source, platform='wechat'))
    run_job(project, 'validation.account_count', {}, project.objects('dataset'))
    result = project.objects('result')[0]
    page = QueryService(project).page(result['id'], sort='platform')
    assert page['total'] == 2
    assert [row['record_count'] for row in page['rows']] == [6, 3]
    sources = QueryService(project).sources(result['id'], page['rows'][0]['result_row_id'])
    assert sources['total'] == 6
    assert len({row['source_id'] for row in sources['rows']}) == 2
    assert sorted(row['source_row'] for row in sources['rows']) == [3, 3, 4, 4, 5, 5]


def test_index_detects_version_changes_and_rebuild(tmp_path):
    source = workbook(tmp_path / 'a.xlsx', [['00123', '1', None, '00045', None]])
    project = LocalProject(tmp_path / 'p')
    query = QueryService(project)
    assert query.search_index('00123')['stale']
    run_job(project, 'xlsx.import', parameters(source))
    run_job(project, 'index.rebuild', {}, project.objects('dataset'))
    assert query.search_index('00123')['total'] == 1
    assert query.search_index('00045')['total'] == 1
    run_job(project, 'xlsx.import', parameters(source))
    assert query.search_index('00123')['stale']
    run_job(project, 'index.rebuild', {}, project.objects('dataset'))
    assert query.search_index('00123')['total'] == 2
    assert not query.search_index('00123')['stale']


def test_formula_cache_numeric_raw_and_shared_strings(tmp_path):
    path = workbook(tmp_path / 'a.xlsx', [['00123', '=0.1+0.2', '2026-10-22', '00045', 'x']])
    assert validate_import(parameters(path))['rejected_count'] == 1
    # A real cached formula numeric payload, created directly to avoid Excel/Python float rewriting.
    with ZipFile(path) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    xml = files['xl/worksheets/sheet1.xml'].decode()
    xml = re.sub(r'(<f>0\.1\+0\.2</f>)<v(?:\s*/>|></v>)', r'\1<v>0.300000000001</v>', xml)
    xml = xml.replace('t="inlineStr"><is><t>00123</t></is>', 't="s"><v>0</v>')
    files['xl/worksheets/sheet1.xml'] = xml.encode()
    files['xl/sharedStrings.xml'] = b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>00123</t></si></sst>'
    with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    with WorkbookReader(path) as reader:
        rows = dict(reader.rows('流水'))
    assert rows[3][1]['value'] == '00123'
    assert rows[3][2]['value'] == '0.300000000001'
    assert rows[3][2]['formula'] == '=0.1+0.2'
    assert validate_import(parameters(path))['valid_count'] == 1


def test_cancel_and_filter_validation_do_not_publish(tmp_path):
    source = workbook(tmp_path / 'a.xlsx', [['00123', '1', None, None, None]])
    with pytest.raises(CancelledError):
        validate_import(parameters(source), cancelled=lambda: True)
    project = LocalProject(tmp_path / 'p')
    run_job(project, 'xlsx.import', parameters(source))
    dataset = project.objects('dataset')[0]
    with pytest.raises(DataError):
        QueryService(project).page(dataset['id'], [{'column': 'amount); DROP TABLE x;--', 'op': 'eq', 'value': 1}])
    with pytest.raises(DataError):
        QueryService(project).page(dataset['id'], [{'column': 'amount', 'op': 'bad', 'value': 1}])
    assert pq.ParquetFile(project.path / dataset['data']['data_path']).metadata.num_rows == 1

def test_date_precision_and_long_numeric_identifiers_are_explicit(tmp_path):
    path = workbook(tmp_path / 'date.xlsx', [
        ['00123', '1', '2026-10-22 12:34', None, None],
        ['00123', '2', '2026-10-22 12:34:56.123456789+08:00', None, None],
        [6222021234567890123, '3', '2026-10-22', None, None],
    ])
    project = LocalProject(tmp_path / 'p')
    run_job(project, 'xlsx.import', parameters(path))
    rows = QueryService(project).page(project.objects('dataset')[0]['id'], sort='source_row')['rows']
    assert rows[0]['occurred_at'] == '2026-10-22T12:34'
    assert rows[0]['time_precision'] == 'minute'
    assert rows[1]['occurred_at'] == '2026-10-22T12:34:56.123456789+08:00'
    assert rows[1]['time_precision'] == 'nanosecond'
    assert rows[1]['timezone'] == '+0800'
    assert 'Excel' in rows[2]['quality_flags']


def test_precise_time_filters_preserve_fraction_and_separate_known_unknown(tmp_path):
    values = [
        '2026-10-22 12:34',
        '2026-10-22T12:34:00.000',
        '2026-10-22T12:34:00.000000001',
        '2026-10-22T12:34:00.1234567890123+08:00',
        '2026-10-22T04:34:00.123456789012300Z',
        '2026-10-22T04:34:00.1234567890124Z',
        '2026-10-22',
        None,
    ]
    source = workbook(tmp_path / 'times.xlsx', [[f'{i:05d}', '1', value, None, None]
                                               for i, value in enumerate(values)])
    project = LocalProject(tmp_path / 'p')
    run_job(project, 'xlsx.import', parameters(source))
    dataset = project.objects('dataset')[0]
    query = QueryService(project)

    def matches(value, op='eq'):
        return query.page(dataset['id'], [{'column': 'occurred_at', 'op': op, 'value': value}],
                          sort='source_row')

    assert [row['owner_account'] for row in matches('2026-10-22 12:34')['rows']] == ['00000', '00001']
    assert [row['owner_account'] for row in matches('2026-10-22T12:34:00.000')['rows']] == ['00000', '00001']
    aware = matches('2026-10-22 04:34:00.1234567890123Z')
    assert [row['owner_account'] for row in aware['rows']] == ['00003', '00004']
    assert aware['rows'][0]['occurred_at'] == values[3]
    assert aware['rows'][1]['occurred_at'] == values[4]
    assert aware['time_comparison']['version'] == 'time-query/1.0.0'
    assert matches('2026-10-22 04:34:00.1234567890123')['total'] == 0
    filters = [
        {'column': 'occurred_at', 'op': 'gte', 'value': '2026-10-22 12:34:00.1234567890123+08:00'},
        {'column': 'occurred_at', 'op': 'lt', 'value': '2026-10-22 04:34:00.1234567890124Z'},
    ]
    page = query.page(dataset['id'], filters, sort='occurred_at', descending=True)
    assert page['total'] == 2
    for format_name in ('csv', 'xlsx'):
        manifest, staging = run_job(project, 'result.export', {
            'object_id': dataset['id'], 'filters': filters, 'scope': 'filtered',
            'format': format_name, 'destination': str(tmp_path / ('time-export.' + format_name)),
            'sort': 'occurred_at', 'descending': True,
        }, [dataset])
        if format_name == 'csv':
            with (staging / 'export.csv').open(encoding='utf-8-sig', newline='') as stream:
                exported = list(csv.DictReader(stream))
        else:
            book = openpyxl.load_workbook(staging / 'export.xlsx', read_only=True)
            headers, *rows = list(book.active.values)
            exported = [dict(zip(headers, row, strict=True)) for row in rows]
            book.close()
        assert [row['record_id'] for row in exported] == [row['record_id'] for row in page['rows']]
        assert [row['occurred_at'] for row in exported] == [row['occurred_at'] for row in page['rows']]
        assert not any(key.startswith('__traceatlas_time_') for key in exported[0])
        assert manifest['summary']['row_count'] == page['total']


def test_time_sort_groups_categories_and_calendar_filters_keep_original_day(tmp_path):
    source = workbook(tmp_path / 'mixed.xlsx', [
        ['late-aware', '1', '2026-10-22T12:00:00+08:00', None, None],
        ['early-aware', '1', '2026-10-22T03:00:00Z', None, None],
        ['date-crossing', '1', '2026-10-23T00:30:00+08:00', None, None],
        ['late-unknown', '1', '2026-10-22T05:00:00', None, None],
        ['early-unknown', '1', '2026-10-22T04:00', None, None],
        ['date-only', '1', '2026-10-22', None, None],
        ['missing', '1', None, None, None],
    ])
    project = LocalProject(tmp_path / 'p')
    run_job(project, 'xlsx.import', parameters(source))
    dataset = project.objects('dataset')[0]
    query = QueryService(project)
    ascending = query.page(dataset['id'], sort='occurred_at')
    assert [row['owner_account'] for row in ascending['rows']] == [
        'early-aware', 'late-aware', 'date-crossing', 'early-unknown', 'late-unknown', 'date-only', 'missing']
    descending = query.page(dataset['id'], sort='occurred_at', descending=True)
    assert [row['owner_account'] for row in descending['rows']] == [
        'date-crossing', 'late-aware', 'early-aware', 'late-unknown', 'early-unknown', 'date-only', 'missing']
    calendar = query.page(dataset['id'], [{'column': 'occurred_at', 'op': 'eq', 'value': '2026-10-22'}])
    assert calendar['total'] == 5
    assert 'date-crossing' not in {row['owner_account'] for row in calendar['rows']}
    assert query.page(dataset['id'], [{'column': 'occurred_at', 'op': 'gt', 'value': '2026-10-22'}])['total'] == 1
    pages = [query.page(dataset['id'], sort='occurred_at', offset=offset, limit=2)['rows'] for offset in (0, 2, 4, 6)]
    assert [row['record_id'] for part in pages for row in part] == [row['record_id'] for row in ascending['rows']]


def test_explicit_beijing_and_unsupported_zone_are_not_treated_as_unknown(tmp_path):
    source = workbook(tmp_path / 'zones.xlsx', [['00123', '1', '2026-10-22 12:34', None, None]])
    project = LocalProject(tmp_path / 'p')
    for zone in ('Asia/Shanghai', 'Unrecognized/Region', 'unknown'):
        run_job(project, 'xlsx.import', parameters(source, timezone=zone))
    query = QueryService(project)
    aware = [{'column': 'occurred_at', 'op': 'eq', 'value': '2026-10-22T04:34Z'}]
    naive = [{'column': 'occurred_at', 'op': 'eq', 'value': '2026-10-22T12:34'}]
    datasets = project.objects('dataset')
    assert query.page(datasets[0]['id'], aware)['total'] == 1
    assert query.page(datasets[0]['id'], naive)['total'] == 0
    assert query.page(datasets[1]['id'], aware)['total'] == 0
    assert query.page(datasets[1]['id'], naive)['total'] == 0
    assert query.page(datasets[1]['id'], sort='occurred_at')['total'] == 1
    assert query.page(datasets[2]['id'], naive)['total'] == 1
    for value in ('2026-02-30 12:00', '2026-10-22 12:00+24:00', '22/10/2026 12:00'):
        with pytest.raises(DataError, match='时间|时区'):
            query.page(datasets[0]['id'], [{'column': 'occurred_at', 'op': 'gte', 'value': value}])


def test_precise_time_comparison_supports_year_range_and_subnanoseconds(tmp_path):
    values = ['0001-01-01T00:00:00.000000000001Z', '9999-12-31T23:59:59.999999999999Z',
              '2026-10-22T12:34:00.00000000000000000000000000000000000001',
              '2026-10-22T12:34:00.00000000000000000000000000000000000002']
    source = workbook(tmp_path / 'ranges.xlsx', [[str(i), '1', value, None, None]
                                                for i, value in enumerate(values)])
    project = LocalProject(tmp_path / 'p')
    run_job(project, 'xlsx.import', parameters(source))
    dataset = project.objects('dataset')[0]
    query = QueryService(project)
    for value in values:
        page = query.page(dataset['id'], [{'column': 'occurred_at', 'op': 'eq', 'value': value}])
        assert page['total'] == 1
        assert page['rows'][0]['occurred_at'] == value
    assert query.page(dataset['id'], [{'column': 'occurred_at', 'op': 'eq', 'value': '9999-12-31'}])['total'] == 1


def test_xlsx_does_not_silently_truncate_long_raw_cells(tmp_path):
    source = workbook(tmp_path / 'long.xlsx', [['00123', '1', None, 'A' * 20000, 'B' * 20000]])
    project = LocalProject(tmp_path / 'p')
    run_job(project, 'xlsx.import', parameters(source))
    dataset = project.objects('dataset')[0]
    with pytest.raises(DataError) as error:
        run_job(project, 'result.export', {'object_id': dataset['id'], 'format': 'xlsx',
                                         'destination': str(tmp_path / 'long-export.xlsx')}, [dataset])
    assert error.value.code == 'EXPORT_CELL_LIMIT'
    assert not project.objects('export')
    manifest, staging = run_job(project, 'result.export', {
        'object_id': dataset['id'], 'format': 'csv', 'destination': str(tmp_path / 'long.csv')}, [dataset])
    with (staging / 'export.csv').open(encoding='utf-8-sig', newline='') as stream:
        row = next(csv.DictReader(stream))
    assert len(row['raw_cells']) > 40000
    assert manifest['summary']['row_count'] == 1


def test_zero_rows_is_a_complete_persisted_result(tmp_path):
    source = workbook(tmp_path / 'empty.xlsx', [])
    project = LocalProject(tmp_path / 'p')
    run_job(project, 'xlsx.import', parameters(source))
    run_job(project, 'validation.account_count', {}, project.objects('dataset'))
    result = project.objects('result')[0]
    assert result['data']['completion'] == 'complete'
    assert QueryService(project).page(result['id'])['total'] == 0

def test_source_hash_change_is_blocked_before_commit(tmp_path):
    source = workbook(tmp_path / 'a.xlsx', [['00123', '1', None, None, None]])
    project = LocalProject(tmp_path / 'p')
    with pytest.raises(DataError) as error:
        run_job(project, 'xlsx.import', parameters(source, expected_source_hash='0' * 64))
    assert error.value.code == 'SOURCE_CHANGED'
    assert not project.objects()


def test_balance_precision_unknown_direction_and_original_status(tmp_path):
    source = workbook(tmp_path / 'a.xlsx', [['00123', '-0.123456789012', None, 'in', '撤销']])
    project = LocalProject(tmp_path / 'p')
    args = parameters(source, mapping={'owner_account': 1, 'amount': 2, 'balance': 3,
                                      'direction': 4, 'status': 5})
    run_job(project, 'xlsx.import', args)
    dataset = project.objects('dataset')[0]
    row = QueryService(project).page(dataset['id'])['rows'][0]
    assert row['amount'] == Decimal('-0.123456789012')
    assert row['balance'] is None
    assert row['status'] == '撤销'
    assert row['direction_raw'] == row['direction'] == 'in'
    assert any('冲突' in issue['message'] for issue in json.loads(row['quality_flags']))
    source = workbook(tmp_path / 'b.xlsx', [['00123', '1', '2.123456789012', '借', '0']])
    args['path'] = str(source)
    run_job(project, 'xlsx.import', args)
    row = QueryService(project).page(project.objects('dataset')[1]['id'])['rows'][0]
    assert row['balance'] == Decimal('2.123456789012')
    assert row['direction'] == 'unknown'
    assert row['direction_raw'] == '借'
    assert row['status'] == '0'
