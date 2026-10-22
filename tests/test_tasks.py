from __future__ import annotations

import csv
import io
import json
import subprocess
import sys
import time

import openpyxl
import pytest

from traceatlas.contracts import TraceAtlasError
from traceatlas.project import Project
from traceatlas.query import QueryService
from traceatlas.tasks import TaskManager


def sample(path, count=100):
    book = openpyxl.Workbook(write_only=True)
    sheet = book.create_sheet('流水')
    sheet.append(['账号', '金额', '时间'])
    for i in range(count):
        sheet.append([f'{i % 12:05d}', '100.123456789012', '2026-10-22'])
    book.save(path)
    return {'path': str(path), 'sheet': '流水', 'header_row': 1,
            'mapping': {'owner_account': 1, 'amount': 2, 'occurred_at': 3},
            'platform': 'bank', 'account_namespace': 'bank-account', 'allow_partial': False}


def wait(manager, tid, timeout=30):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        task = manager.project.get_task(tid)
        if task['state'] not in {'queued', 'running', 'cancel_requested', 'committing'}:
            return task
        time.sleep(0.05)
    raise AssertionError(f'task timed out: {manager.project.get_task(tid)}')


def test_real_worker_roundtrip_reopen_and_export(tmp_path):
    params = sample(tmp_path / '来源 流水.xlsx')
    project = Project.create(tmp_path / '审计项目 01', '审计项目 01')
    manager = TaskManager(project)
    try:
        task = wait(manager, manager.submit('xlsx.import', params))
        assert task['state'] == 'succeeded', task
        dataset = project.objects('dataset')[0]
        task = wait(manager, manager.submit('validation.account_count', {}, [dataset['id']]))
        assert task['state'] == 'succeeded', task
        result = project.objects('result')[0]
        page = QueryService(project).page(result['id'])
        assert page['total'] == 12
        assert sum(row['record_count'] for row in page['rows']) == 100
        row = page['rows'][0]
        assert QueryService(project).sources(result['id'], row['result_row_id'])['total'] == row['record_count']
        export = tmp_path / '结果.csv'
        task = wait(manager, manager.submit('result.export', {'object_id': result['id'], 'scope': 'all',
                        'format': 'csv', 'destination': str(export), 'overwrite': False}))
        assert task['state'] == 'succeeded', task
        with export.open(encoding='utf-8-sig', newline='') as stream:
            assert len(list(csv.reader(stream))) == 13
        project.save_session({'selected_object_id': result['id']})
        project_id = project.id
    finally:
        manager.shutdown(cancel=True)
        project.close()
    reopened = Project.open(tmp_path / '审计项目 01')
    assert reopened.id == project_id
    assert reopened.load_session()['selected_object_id'] == result['id']
    assert QueryService(reopened).page(result['id'])['total'] == 12
    assert len(reopened.list_tasks()) == 3
    reopened.close()


def test_queue_cancel_and_worker_kill_retry(tmp_path):
    params = sample(tmp_path / '长任务.xlsx', 20000)
    project = Project.create(tmp_path / 'project', 'project')
    manager = TaskManager(project)
    try:
        tid = manager.submit('xlsx.import', params)
        queued = manager.submit('xlsx.import', params)
        manager.cancel(queued)
        assert project.get_task(queued)['state'] == 'cancelled'
        deadline = time.monotonic() + 10
        while not manager._process and time.monotonic() < deadline:
            time.sleep(0.02)
        assert manager._process
        manager._process.kill()
        task = wait(manager, tid)
        assert task['state'] == 'interrupted', task
        assert project.objects('dataset') == []
        old_attempt = task['current_attempt_id']
        manager.retry(tid)
        task = wait(manager, tid, timeout=60)
        assert task['state'] == 'succeeded', task
        assert task['current_attempt_id'] != old_attempt
        assert len(project.objects('dataset')) == 1
        assert len(project.task_attempts(tid)) == 2
        assert project.objects('dataset')[0]['data']['row_count'] == 20000
    finally:
        manager.shutdown(cancel=True)
        project.close()


def test_invalid_input_error_does_not_create_success(tmp_path):
    params = sample(tmp_path / 'bad.xlsx', 3)
    book = openpyxl.load_workbook(params['path'])
    book.active['B3'] = '非法金额'
    book.save(params['path'])
    project = Project.create(tmp_path / 'project', 'project')
    manager = TaskManager(project)
    try:
        tid = manager.submit('xlsx.import', params)
        task = wait(manager, tid)
        assert task['state'] == 'failed'
        assert task['error']['code']
        assert project.objects('dataset') == []
        with pytest.raises(TraceAtlasError, match="处理输入问题"):
            manager.retry(tid)
        params['allow_partial'] = True
        task = wait(manager, manager.submit('xlsx.import', params))
        assert task['state'] == 'succeeded', task
        assert project.objects('dataset')[0]['data']['row_count'] == 2
        assert project.objects('dataset')[0]['data']['rejected_count'] == 1
        assert json.dumps(task['summary'], ensure_ascii=False)
    finally:
        manager.shutdown(cancel=True)
        project.close()


@pytest.mark.parametrize(('changed', 'expected'), [
    ({'module_version': '0.9.0'}, 'TASK_VERSION_UNSUPPORTED'),
    ({'parameter_schema_version': 2}, 'TASK_SCHEMA_UNSUPPORTED'),
    ({'protocol_version': 2}, 'TASK_SCHEMA_UNSUPPORTED'),
])
def test_retry_refuses_incompatible_saved_task_without_new_attempt(tmp_path, changed, expected):
    with Project.create(tmp_path / 'project', 'version checks') as project:
        manager = TaskManager(project)
        try:
            spec = {'feature_id': 'xlsx.import', 'module_version': '1.0.0', 'protocol_version': 1,
                    'parameter_schema_version': 1, 'parameters': sample(tmp_path / 'source.xlsx', 1)}
            spec.update(changed)
            task = project.create_task(spec)
            project.update_task(task['task_id'], 'interrupted', error={'retryable': True})
            with pytest.raises(TraceAtlasError) as error:
                manager.retry(task['task_id'])
            assert error.value.code == expected
            assert len(project.task_attempts(task['task_id'])) == 1
            assert project.get_task(task['task_id'])['state'] == 'interrupted'
            assert project.objects() == []
        finally:
            manager.shutdown(cancel=True)


def test_queued_task_rechecks_version_before_starting_worker(tmp_path):
    with Project.create(tmp_path / 'project', 'queued version check') as project:
        manager = TaskManager(project)
        try:
            with manager._lock:
                tid = manager.submit('xlsx.import', sample(tmp_path / 'source.xlsx', 1))
                assert project.get_task(tid)['parameter_schema_version'] == 1
                manager.registry._features['xlsx.import']['module_version'] = '2.0.0'
            task = wait(manager, tid)
            assert task['state'] == 'failed'
            assert task['error']['code'] == 'TASK_VERSION_UNSUPPORTED'
            assert 'worker_pid' not in task
            assert project.objects() == []
        finally:
            manager.shutdown(cancel=True)


def test_worker_rejects_incompatible_parameter_format_itself(tmp_path):
    staging = tmp_path / 'staging'
    staging.mkdir()
    spec = {'task_id': 'task', 'attempt_id': 'attempt', 'feature_id': 'xlsx.import',
            'module_version': '1.0.0', 'protocol_version': 1, 'parameter_schema_version': 2,
            'parameters': {}, 'staging_path': str(staging)}
    spec_path = staging / 'task.json'
    spec_path.write_text(json.dumps(spec), encoding='utf-8')
    completed = subprocess.run([sys.executable, '-m', 'traceatlas.worker', str(spec_path)],
                               check=False, input='', capture_output=True, text=True, timeout=20)
    events = [json.loads(line) for line in completed.stdout.splitlines()]
    assert completed.returncode == 1
    assert events[-1]['error']['code'] == 'TASK_SCHEMA_UNSUPPORTED'
    assert not (staging / 'manifest.json').exists()


def test_worker_memory_error_is_retryable_and_not_a_success(tmp_path, monkeypatch, capsys):
    from traceatlas.worker import main

    spec = {'task_id': 'task', 'attempt_id': 'attempt', 'feature_id': 'xlsx.import',
            'module_version': '1.0.0', 'protocol_version': 1, 'parameter_schema_version': 1,
            'parameters': {'path': 'source.xlsx', 'sheet': '流水', 'mapping': {'amount': 1}},
            'staging_path': str(tmp_path)}
    spec_path = tmp_path / 'task.json'
    spec_path.write_text(json.dumps(spec), encoding='utf-8')

    def out_of_memory(*_args):
        raise MemoryError

    monkeypatch.setattr('traceatlas.jobs.execute_job', out_of_memory)
    monkeypatch.setattr(sys, 'stdin', io.StringIO(''))
    assert main(str(spec_path)) == 1
    events = [json.loads(line) for line in capsys.readouterr().out.splitlines()]
    assert events[-1]['error']['code'] == 'MEMORY_LIMIT'
    assert events[-1]['error']['retryable'] is True
    assert not (tmp_path / 'manifest.json').exists()
