"""Desktop integration tests use real project storage, worker jobs and Qt events."""

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from PySide6.QtCore import QSettings, Qt
from PySide6.QtWidgets import QDialog

from traceatlas.contracts import TraceAtlasError
from traceatlas.project import Project
from traceatlas.ui.common import copy_error_report, destination_snapshot
from traceatlas.ui.dialogs import ExportDialog, ImportDialog
from traceatlas.ui.main_window import MainWindow
from traceatlas.ui.sources import SourcesDialog


@pytest.fixture(autouse=True)
def isolated_settings(tmp_path):
    QSettings.setDefaultFormat(QSettings.Format.IniFormat)
    QSettings.setPath(QSettings.Format.IniFormat, QSettings.Scope.UserScope, str(tmp_path / "settings"))


@pytest.fixture
def workbook(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "流水明细"
    sheet.append(["本方账户", "交易金额", "对方账户", "交易日期"])
    sheet.append(["00123", "100.125", "98765", "2026-10-01"])
    sheet.append(["00123", "20.25", "98766", "2026-10-02"])
    sheet.append(["00123", "30.000000000001", "98767", "2026-10-03"])
    sheet.append(["00999", "40.00", "98767", "2026-10-04"])
    sheet.append(["00123", "非法金额", "98768", "2026-10-05"])
    path = tmp_path / "审计 流水.xlsx"
    book.save(path)
    return path


def wait_finished(qtbot, window, task_id):
    qtbot.waitUntil(
        lambda: (
            window.project.get_task(task_id)["state"] in {"succeeded", "failed", "cancelled", "interrupted"}
        ),
        timeout=30000,
    )
    task = window.project.get_task(task_id)
    assert task["state"] == "succeeded", task
    window.poll_tasks()


def test_import_requires_quality_review_and_explicit_partial(qtbot, workbook):
    dialog = ImportDialog(path=str(workbook))
    qtbot.addWidget(dialog)
    dialog.show()
    qtbot.waitUntil(lambda: dialog.preview is not None, timeout=15000)
    assert dialog.mapping["owner_account"].currentData() == 1
    assert not dialog.accept_button.isEnabled()
    qtbot.mouseClick(dialog.validate_button, Qt.MouseButton.LeftButton)
    qtbot.waitUntil(lambda: dialog.report is not None, timeout=15000)
    assert dialog.report["input_count"] == 5
    assert dialog.report["valid_count"] == 4
    assert dialog.report["rejected_count"] == 1
    assert dialog.report_model.rows[0]["row"] == 6
    assert not dialog.accept_button.isEnabled()
    qtbot.mouseClick(dialog.partial, Qt.MouseButton.LeftButton)
    assert dialog.accept_button.isEnabled()
    dialog.start.setValue(3)
    assert not dialog.accept_button.isEnabled(), "Changing input scope invalidates the reviewed report"
    dialog.reject()


def test_desktop_vertical_chain_query_source_export_reopen(qtbot, workbook, tmp_path, monkeypatch):
    project = Project.create(tmp_path / "审计项目 01", "审计项目 01")
    window = MainWindow(project)
    qtbot.addWidget(window)
    window.show()
    parameters = {
        "path": str(workbook),
        "sheet": "流水明细",
        "header_row": 1,
        "start_row": 2,
        "mapping": {"owner_account": 1, "amount": 2, "counterparty_account": 3, "occurred_at": 4},
        "platform": "bank",
        "account_namespace": "account",
        "currency": "CNY",
        "timezone": "unknown",
        "allow_partial": True,
    }
    task = window.submit_task("xlsx.import", parameters)
    assert task
    # Navigation remains usable while the worker imports.
    window.workspace.setCurrentWidget(window.modules_page)
    wait_finished(qtbot, window, task)
    dataset = project.objects("dataset")[0]
    window.show_object(dataset["id"])
    qtbot.waitUntil(lambda: window._query_ready, timeout=15000)
    assert window.total == 4
    window.filter_column.setCurrentIndex(window.filter_column.findData("owner_account"))
    window.filter_operator.setCurrentIndex(window.filter_operator.findData("eq"))
    window.filter_value.setText("00123")
    qtbot.mouseClick(window.filter_button, Qt.MouseButton.LeftButton)
    qtbot.waitUntil(lambda: window._query_ready, timeout=15000)
    assert window.total == 3
    assert {row["owner_account"] for row in window.table_model.rows} == {"00123"}
    destination = tmp_path / "筛选导出.xlsx"
    export = window.submit_task(
        "result.export",
        {
            "object_id": dataset["id"],
            "filters": window.filters,
            "sort": "source_row",
            "descending": False,
            "scope": "filtered",
            "format": "xlsx",
            "destination": str(destination),
            "overwrite": False,
        },
        [dataset["id"]],
    )
    wait_finished(qtbot, window, export)
    assert destination.is_file()
    with destination.open("rb") as stream:
        book = load_workbook(stream, read_only=True)
        assert book.active.max_row in (None, 4)
        rows = list(book.active.values)
        assert len(rows) == 4
        assert any("00123" in row for row in rows[1:])
        book.close()
    counted = window.submit_task("validation.account_count", {}, [dataset["id"]])
    wait_finished(qtbot, window, counted)
    result = project.objects("result")[0]
    window.show_object(result["id"])
    qtbot.waitUntil(lambda: window._query_ready, timeout=15000)
    row = next(value for value in window.table_model.rows if value["owner_account"] == "00123")
    assert row["record_count"] == 3
    source_dialog = SourcesDialog(project, result["id"], row["result_row_id"], window)
    qtbot.addWidget(source_dialog)
    source_dialog.show()
    qtbot.waitUntil(
        lambda: source_dialog.total == 3 and bool(source_dialog.original_model.rows), timeout=15000
    )
    assert {value["source_row"] for value in source_dialog.contributions_model.rows} == {2, 3, 4}
    assert source_dialog.original.selectionModel().selectedRows()
    qtbot.waitUntil(lambda: not source_dialog._jobs, timeout=15000)
    source_dialog.reject()
    # A missing original does not hide the result; the UI can locate the same source again.
    stored_source = project.path / project.objects("source")[0]["data"]["stored_path"]
    stored_source.unlink()
    monkeypatch.setattr("traceatlas.ui.sources.show_error", lambda *_: None)
    monkeypatch.setattr("traceatlas.ui.sources.QFileDialog.getOpenFileName", lambda *_: (str(workbook), ""))
    repair = SourcesDialog(project, result["id"], row["result_row_id"], window)
    qtbot.addWidget(repair)
    repair.show()
    qtbot.waitUntil(lambda: repair.repair_button.isEnabled(), timeout=15000)
    qtbot.mouseClick(repair.repair_button, Qt.MouseButton.LeftButton)
    qtbot.waitUntil(
        lambda: stored_source.exists() and bool(repair.original_model.rows) and not repair._jobs,
        timeout=15000,
    )
    repair.reject()
    saved_path = project.path
    window.save_project()
    window.close()
    reopened = Project.open(saved_path)
    second = MainWindow(reopened)
    qtbot.addWidget(second)
    second.show()
    qtbot.waitUntil(lambda: second._query_ready, timeout=15000)
    assert second.current_object["id"] == result["id"]
    assert second.total == 2
    assert second.tasks_table.topLevelItemCount() == 3
    second.close()


def test_readonly_project_disables_mutations(qtbot, tmp_path):
    writable = Project.create(tmp_path / "共享项目", "共享项目")
    readonly = Project.open(writable.path, readonly=True)
    window = MainWindow(readonly)
    qtbot.addWidget(window)
    window.show()
    assert "只读" in window.windowTitle()
    assert not window.import_action.isEnabled()
    assert not window.save_action.isEnabled()
    assert not window.index_rebuild_button.isEnabled()
    assert window.manager is None
    window.close()
    writable.close()


def test_export_confirmation_describes_scope_encoding_and_split(qtbot, tmp_path):
    dialog = ExportDialog("审计结果", 12, 1_100_000, tmp_path)
    qtbot.addWidget(dialog)
    assert "12" in dialog.scope.currentText()
    dialog.scope.setCurrentIndex(1)
    assert "多个工作表" in dialog.details.text()
    dialog.format.setCurrentIndex(1)
    assert "UTF-8" in dialog.details.text()
    assert "BOM" in dialog.details.text()
    assert Path(dialog.destination.text()).suffix == ".csv"
    dialog.accept()
    assert dialog.result() == QDialog.DialogCode.Accepted
    assert dialog.parameters()["scope"] == "all"
    assert dialog.parameters()["overwrite"] is False


def test_quality_error_copy_protects_project_and_publishes_complete_file(tmp_path):
    project_root = tmp_path / "项目"
    (project_root / "exports").mkdir(parents=True)
    (project_root / "datasets").mkdir()
    source = project_root / "datasets" / "errors.jsonl"
    source.write_text('{"row":6,"message":"金额错误"}\n', encoding="utf-8")
    protected = project_root / "project.json"
    protected.write_bytes(b"project metadata")
    with pytest.raises(TraceAtlasError) as caught:
        copy_error_report(source, protected, project_root, destination_snapshot(protected))
    assert caught.value.code == "INVALID_EXPORT_DESTINATION"
    assert protected.read_bytes() == b"project metadata"
    destination = project_root / "exports" / "完整错误清单.jsonl"
    copy_error_report(source, destination, project_root, None)
    assert destination.read_bytes() == source.read_bytes()
    external = tmp_path / "旧错误清单.jsonl"
    external.write_bytes(b"previous report")
    copy_error_report(source, external, project_root, destination_snapshot(external))
    assert external.read_bytes() == source.read_bytes()
    assert not list(tmp_path.rglob("*.tmp"))


@pytest.mark.parametrize("existing", [False, True])
def test_quality_error_copy_failure_preserves_target_and_cleans_temp(tmp_path, monkeypatch, existing):
    from traceatlas.ui import common

    project_root = tmp_path / "项目"
    project_root.mkdir()
    source = project_root / "errors.jsonl"
    source.write_bytes(b"complete report")
    destination = tmp_path / "保存结果.jsonl"
    if existing:
        destination.write_bytes(b"previous report")
    expected = destination_snapshot(destination)

    def disk_full(reader, writer, length):
        writer.write(b"partial")
        raise OSError(28, "No space left on device")

    monkeypatch.setattr(common.shutil, "copyfileobj", disk_full)
    with pytest.raises(OSError, match="No space"):
        copy_error_report(source, destination, project_root, expected)
    if existing:
        assert destination.read_bytes() == b"previous report"
    else:
        assert not destination.exists()
    assert not list(tmp_path.glob("*.tmp"))


@pytest.mark.parametrize("existing", [False, True])
def test_quality_error_copy_rejects_concurrent_target_change(tmp_path, monkeypatch, existing):
    from traceatlas.ui import common

    project_root = tmp_path / "项目"
    project_root.mkdir()
    source = project_root / "errors.jsonl"
    source.write_bytes(b"complete report")
    destination = tmp_path / "保存结果.jsonl"
    if existing:
        destination.write_bytes(b"previous report")
    expected = destination_snapshot(destination)
    original_copy = common.shutil.copyfileobj

    def changed_target(reader, writer, length):
        original_copy(reader, writer, length)
        destination.write_bytes(b"other application changed this file")

    monkeypatch.setattr(common.shutil, "copyfileobj", changed_target)
    with pytest.raises(TraceAtlasError) as caught:
        copy_error_report(source, destination, project_root, expected)
    assert caught.value.code == "EXPORT_CONFLICT"
    assert destination.read_bytes() == b"other application changed this file"
    assert not list(tmp_path.glob("*.tmp"))
